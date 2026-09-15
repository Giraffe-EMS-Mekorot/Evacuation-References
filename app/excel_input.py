"""Reads a manually-filled Excel file (NOT a scanned/photographed
certificate run through Claude Vision) directly into record dicts, for the
"someone already typed the data into a tracking sheet" case - see
streamlit_app.py's upload area, which auto-detects a .xlsx by extension and
routes it here instead of through app.extractor.extract_certificate_pages.

**Deliberately bypasses this whole codebase's per-record AI-extraction
safety nets** - NameNormalizer (app/normalize.py), the quantity-history
outlier check and the gross/tare/net mismatch check (both effectively N/A
here anyway, but still), closed-list enforcement, and the date-plausibility
check. All of those exist to catch a VISION MODEL's mistakes (a misread
digit, a name spelled differently certificate to certificate, a garbled
OCR read); a manually-typed row was never extracted by one, so running it
through them would either do nothing useful or actively mangle input this
module is specifically told to keep verbatim (see quantity handling below).
Records built here never go through app.pipeline.process_files at all -
streamlit_app.py's upload handler appends them straight into
st.session_state.records, alongside (not instead of) whatever
process_files() returned for any PDF/image files uploaded in the same
batch.

Confidence here defaults to "גבוהה" (a human already reviewed this data -
there's no reason to flag it "for manual review" the way a shaky vision
read is flagged) except for the one gap this feature spec explicitly calls
out: a missing year, which genuinely needs manual follow-up and must not be
silently guessed from context (e.g. the sheet's own tab/file name) - see
_MISSING_YEAR_NOTE.
"""
from pathlib import Path
from typing import List

from openpyxl import load_workbook

from .fields import EXTRACTED_CONFIDENCE_KEY, HEBREW_MONTHS, empty_record

# Column-header aliases this reader recognizes, mapped to the record key
# they fill. Supports several real-world header spellings for the same
# field on purpose (the whole point of "recognize columns by name, not
# fixed position") - both this pipeline's own current output headers (so a
# previously-produced ריכוז_תעודות.xlsx can be re-uploaded and re-read) and
# the header names a hand-built tracking sheet is likely to use instead.
# _year/_month are internal sentinel keys (not real FIELD_DEFS keys on
# their own) combined into a single "date" field below - see
# _build_date_field(); everything else maps straight to a real record key.
#
# Any column in the uploaded file whose header ISN'T listed here is
# silently ignored - this is what makes "helper" columns sitting off to the
# side (dropdown source lists like "אגפים"/"סוגי פסולת"/"חודשים"/"יחידות")
# harmless without any special-case skip-list code: they simply never match
# a known header, so their values are never read.
_HEADER_TO_KEY = {
    "מרחב": "region",
    "אתר/יחידה": "site",
    "אתר": "site",
    "יחידה": "site",
    "אתר/מקור": "site",
    "שנה": "_year",
    "חודש": "_month",
    "סוג הפסולת": "waste_type",
    "כמות": "quantity",
    "כמות (נטו)": "quantity",
    "יחידת מידה": "unit",
    "אתר קולט": "supplier_or_carrier",
    "מקור אסמכתא": "supplier_or_carrier",
    "סוג אסמכתא": "supplier_or_carrier",
    "ספק/מוביל": "supplier_or_carrier",
    "מספר תעודה/אסמכתא": "certificate_or_reference",
    "הערות": "notes",
}

_MISSING_YEAR_NOTE = "שנה חסרה במקור - נדרשת השלמה ידנית"


def _cell_text(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _build_date_field(year_raw: str, month_raw: str) -> str:
    """Combines a manual sheet's separate year/month cells into this
    pipeline's single "date" field, in the same MM/YYYY shape
    excel_writer.py's _LEGACY_YEAR_KEY/_LEGACY_MONTH_KEY handling already
    reconstructs for a pre-2026-08-30 file (see
    excel_writer.read_existing_records) - reusing that existing shape
    rather than inventing a new one.

    Only forms a clean MM/YYYY when the month cell is EXACTLY one of the
    closed HEBREW_MONTHS names and a year is present. A messier real-world
    value spanning several months (e.g. "אפריל, יוני, ספטמבר נובמבר-2025")
    is deliberately left as raw, uncleaned text instead - per this module's
    whole "accept obviously-nonstandard-but-clearly-intentional human input
    as-is, don't parse/clean it" policy (see read_manual_excel's docstring).
    Trying to force that kind of value into MM/YYYY would silently drop the
    very detail (multiple months) that made it non-standard in the first
    place.
    """
    year_raw, month_raw = (year_raw or "").strip(), (month_raw or "").strip()
    if year_raw and month_raw in HEBREW_MONTHS:
        month_num = HEBREW_MONTHS.index(month_raw) + 1
        return f"{month_num:02d}/{year_raw}"
    return " ".join(part for part in (month_raw, year_raw) if part)


def read_manual_excel(path: Path) -> List[dict]:
    """Reads one manually-filled .xlsx (first sheet only) into a list of
    record dicts shaped like every other record this pipeline produces -
    ready to append straight into st.session_state.records (see this
    module's docstring for why that skips app.pipeline.process_files
    entirely).

    Column matching is by header TEXT on row 1 (see _HEADER_TO_KEY), not
    fixed position, so column order in the hand-built file doesn't matter
    and an unrecognized "helper" column is automatically ignored - no
    special-case code needed for that. A row where every recognized column
    is blank is skipped entirely (not turned into an empty record); every
    other row becomes exactly one record, however sparse.

    Quantity is kept EXACTLY as the cell holds it - a numeric cell stays
    numeric, a human annotation like "כ-1.8" stays that exact text - never
    blanked or rejected for "not parsing as a clean number" the way the
    AI-extraction path (see streamlit_app.py's post-process_files quantity
    block) treats an unparseable value. That's the right call for a vision
    model that might have misread a digit; it's the wrong call for a person
    who wrote "כ-1.8" on purpose. excel_writer.write_records()'s own
    parse_quantity()-with-raw-string-fallback already renders this
    correctly as a plain text cell with no extra code needed here.
    """
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]

    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
    col_keys = [_HEADER_TO_KEY.get(_cell_text(label)) for label in header_row]
    if not any(col_keys):
        return []

    records = []
    for row in ws.iter_rows(min_row=2, max_col=len(col_keys), values_only=True):
        if all(value is None for value in row):
            continue
        raw = {}
        for key, value in zip(col_keys, row):
            if key is None:
                continue  # an unrecognized ("helper") column - never read
            raw[key] = value

        if not any(_cell_text(v) for v in raw.values()):
            continue  # every recognized column blank on this row - nothing to record

        record = empty_record(filename=path.name)
        record["region"] = _cell_text(raw.get("region"))
        record["site"] = _cell_text(raw.get("site"))
        record["waste_type"] = _cell_text(raw.get("waste_type"))
        record["unit"] = _cell_text(raw.get("unit"))
        record["supplier_or_carrier"] = _cell_text(raw.get("supplier_or_carrier"))
        record["certificate_or_reference"] = _cell_text(raw.get("certificate_or_reference"))

        qty_raw = raw.get("quantity")
        record["quantity"] = qty_raw if isinstance(qty_raw, (int, float)) else _cell_text(qty_raw)

        year_raw = _cell_text(raw.get("_year"))
        month_raw = _cell_text(raw.get("_month"))
        record["year"] = year_raw
        record["month"] = month_raw
        record["date"] = _build_date_field(year_raw, month_raw)

        notes = [n for n in (_cell_text(raw.get("notes")),) if n]
        if not year_raw:
            record["confidence"] = "נמוכה"
            notes.append(_MISSING_YEAR_NOTE)
        else:
            record["confidence"] = "גבוהה"
        record["notes"] = " | ".join(notes)

        record["source_file"] = path.name
        # Freeze the confidence this reader assigned, so a later edit to the
        # main table's confidence column can't rewrite the "מעקב פנימי"
        # sheet's audit value - same contract pipeline.process_files applies
        # to AI-extracted rows. See fields.EXTRACTED_CONFIDENCE_KEY.
        record[EXTRACTED_CONFIDENCE_KEY] = record["confidence"]
        records.append(record)
    return records
